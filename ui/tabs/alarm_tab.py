from __future__ import annotations

from pathlib import Path
import re
from dataclasses import replace
from datetime import datetime
from typing import Union

from PyQt5.QtCore import (
    QAbstractTableModel,
    QModelIndex,
    QPoint,
    Qt,
    QObject,
    QRunnable,
    QSortFilterProxyModel,
    QThreadPool,
    pyqtSignal,
)

from PyQt5.QtGui import (
    QColor,
    QPen,
    QPolygon,
)

from PyQt5.QtWidgets import (
    QApplication,
    QAbstractItemView,
    QComboBox,
    QDialog,
    QFileDialog,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QPushButton,
    QLineEdit,
    QProgressDialog,
    QStyledItemDelegate,
    QTableWidget,
    QTableWidgetItem,
    QTableView,
    QVBoxLayout,
    QWidget,
)

from services.alarm_engineers import (
    load_engineers,
)

from workers.alarm_tracking_save_worker import (
    AlarmTrackingSaveWorker,
)


# ============================================================
# COLUMN
# ============================================================

COL_NO = 0
COL_DATE = 1
COL_MACHINE = 2
COL_SLOT = 3
COL_REASON = 4
COL_START = 5
COL_END = 6
COL_TOTAL_TEST = 7
COL_PASS = 8
COL_FAIL = 9
COL_YIELD = 10
COL_LAST15 = 11
COL_TARGET15 = 12
COL_LAST30 = 13
COL_TARGET30 = 14
COL_SCRAP = 15
COL_MODEL = 16
COL_STATUS = 17
COL_DATE_COMPLETE = 18
COL_QUICK_CHECK = 19
COL_CAL_CHECK = 20
COL_ENGINEER_ACTION = 21
COL_MONITOR_DAY1 = 22
COL_MONITOR_DAY2 = 23
COL_MONITOR_DAY3 = 24
COL_COMMENT = 25


# ============================================================
# EDITABLE COLUMNS
# ============================================================

EDITABLE_COLUMNS = (
    COL_STATUS,
    COL_QUICK_CHECK,
    COL_CAL_CHECK,
    COL_ENGINEER_ACTION,
    COL_COMMENT,
)


# ============================================================
# VALID STATUS
# ============================================================

VALID_STATUSES = (
    "Chưa tiến hành",
    "Đang tiến hành",
    "Đã hoàn thành",
)


# ============================================================
# CHOICE / EDIT DELEGATE
# ============================================================

class AlarmChoiceDelegate(QStyledItemDelegate):

    def __init__(self, tab):
        super().__init__(tab.table)
        self.tab = tab

    def createEditor(self, parent, option, index):

        if not index.isValid():
            return None

        if self.tab._alarm_loading:
            return None

        if self.tab.is_saving():
            return None

        # ====================================================
        # STATUS
        # ====================================================

        if index.column() == COL_STATUS:

            editor = QComboBox(parent)

            editor.setEditable(False)

            editor.addItems([
                "Chưa tiến hành",
                "Đang tiến hành",
                "Đã hoàn thành",
            ])

            current = str(
                index.data(Qt.EditRole) or ""
            ).strip()

            if current:
                pos = editor.findText(
                    current,
                    Qt.MatchExactly
                )

                if pos >= 0:
                    editor.setCurrentIndex(pos)

            # ------------------------------------------------
            # Khi chọn xong -> lưu giá trị
            # ------------------------------------------------

            editor.currentIndexChanged.connect(
                lambda _index: self._commit(editor)
            )

            return editor

        # ====================================================
        # QUICK CHECK
        # ====================================================

        if index.column() == COL_QUICK_CHECK:

            editor = QComboBox(parent)

            editor.setEditable(False)

            editor.addItems([
                "",
                "PASS",
                "FAIL",
            ])

            current = str(
                index.data(Qt.EditRole) or ""
            ).strip().upper()

            pos = editor.findText(
                current,
                Qt.MatchExactly
            )

            if pos >= 0:
                editor.setCurrentIndex(pos)

            editor.currentIndexChanged.connect(
                lambda _index: self._commit(editor)
            )

            return editor

        # ====================================================
        # CAL CHECK
        # ====================================================

        if index.column() == COL_CAL_CHECK:

            editor = QComboBox(parent)

            editor.setEditable(False)

            editor.addItems([
                "",
                "PASS",
                "FAIL",
            ])

            current = str(
                index.data(Qt.EditRole) or ""
            ).strip().upper()

            pos = editor.findText(
                current,
                Qt.MatchExactly
            )

            if pos >= 0:
                editor.setCurrentIndex(pos)

            editor.currentIndexChanged.connect(
                lambda _index: self._commit(editor)
            )

            return editor

        # ====================================================
        # ENGINEER ACTION
        # ====================================================

        if index.column() == COL_ENGINEER_ACTION:

            editor = QComboBox(parent)

            editor.setEditable(False)

            engineers = list(
                self.tab._engineers or []
            )

            # -----------------------------------------------
            # Cho phép chọn tên Engineer
            # -----------------------------------------------

            editor.addItem("")

            for engineer in engineers:

                name = str(
                    engineer or ""
                ).strip()

                if name and (
                    editor.findText(
                        name,
                        Qt.MatchExactly
                    ) < 0
                ):

                    editor.addItem(name)

            current = str(
                index.data(Qt.EditRole) or ""
            ).strip()

            # Nếu dữ liệu cũ không nằm trong danh sách
            if current:

                pos = editor.findText(
                    current,
                    Qt.MatchExactly
                )

                if pos < 0:

                    editor.addItem(current)

                    pos = (
                        editor.count() - 1
                    )

                editor.setCurrentIndex(pos)

            editor.currentIndexChanged.connect(
                lambda _index: self._commit(editor)
            )

            return editor

        # ====================================================
        # COMMENT
        # ====================================================

        if index.column() == COL_COMMENT:

            editor = QLineEdit(parent)

            editor.setText(
                str(
                    index.data(Qt.EditRole)
                    or ""
                )
            )

            editor.selectAll()

            editor.editingFinished.connect(
                lambda: self._commit(editor)
            )

            return editor

        return None

    # ========================================================
    # SET EDITOR DATA
    # ========================================================

    def setEditorData(self, editor, index):

        value = str(
            index.data(Qt.EditRole)
            or ""
        )

        if isinstance(editor, QComboBox):

            pos = editor.findText(
                value,
                Qt.MatchExactly
            )

            if pos >= 0:
                editor.setCurrentIndex(pos)

        elif isinstance(editor, QLineEdit):

            editor.setText(value)

    # ========================================================
    # SET MODEL DATA
    # ========================================================

    def setModelData(
        self,
        editor,
        model,
        index,
    ):

        if isinstance(
            editor,
            QComboBox
        ):

            value = editor.currentText()

        elif isinstance(
            editor,
            QLineEdit
        ):

            value = editor.text()

        else:

            return

        model.setData(
            index,
            value,
            Qt.EditRole,
        )

    # ========================================================
    # COMMIT
    # ========================================================

    def _commit(self, editor):

        if editor is None:
            return

        try:

            self.commitData.emit(
                editor
            )

            self.closeEditor.emit(
                editor,
                QStyledItemDelegate.NoHint,
            )

        except Exception as error:

            print(
                "Alarm editor commit error:",
                error,
            )


# ============================================================
# EXPORT WORKER
# ============================================================

class AlarmExportSignals(QObject):

    finished = pyqtSignal(
        str,
        str,
    )


class AlarmExportWorker(QRunnable):

    def __init__(
        self,
        output_path,
        headers,
        rows,
    ):

        super().__init__()

        self.output_path = output_path
        self.headers = headers
        self.rows = rows

        self.signals = AlarmExportSignals()

    def run(self):

        workbook = None
        error_message = ""

        try:

            from openpyxl import Workbook

            from openpyxl.styles import (
                Alignment,
                Font,
                PatternFill,
            )

            from openpyxl.utils import (
                get_column_letter,
            )

            workbook = Workbook()

            sheet = workbook.active

            sheet.title = "List of Alarm"

            # ------------------------------------------------
            # HEADER
            # ------------------------------------------------

            sheet.append(
                self.headers
            )

            # ------------------------------------------------
            # DATA
            # ------------------------------------------------

            for values in self.rows:

                sheet.append(
                    values
                )

                for cell in sheet[
                    sheet.max_row
                ]:

                    if isinstance(
                        cell.value,
                        str,
                    ):

                        cell.data_type = "s"

            # ------------------------------------------------
            # HEADER STYLE
            # ------------------------------------------------

            header_fill = PatternFill(
                fill_type="solid",
                fgColor="E2E8F0",
            )

            for cell in sheet[1]:

                cell.font = Font(
                    bold=True,
                    color="1E293B",
                )

                cell.fill = (
                    header_fill
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            sheet.row_dimensions[
                1
            ].height = 26

            # ------------------------------------------------
            # BODY ALIGNMENT
            # ------------------------------------------------

            for row in sheet.iter_rows(
                min_row=2
            ):

                for cell in row:

                    cell.alignment = Alignment(
                        vertical="center",
                    )

            # ------------------------------------------------
            # FILTER
            # ------------------------------------------------

            last_column = (
                get_column_letter(
                    len(self.headers)
                )
            )

            sheet.auto_filter.ref = (
                f"A1:{last_column}"
                f"{sheet.max_row}"
            )

            sheet.freeze_panes = "A2"

            # ------------------------------------------------
            # WIDTH
            # ------------------------------------------------

            widths = (
                14,
                18,
                18,
                12,
                10,
                23,
                23,
                12,
                30,
                30,
                20,
                24,
                24,
            )

            for column, width in enumerate(
                widths,
                start=1,
            ):

                sheet.column_dimensions[
                    get_column_letter(
                        column
                    )
                ].width = width

            workbook.save(
                self.output_path
            )

        except Exception as error:

            error_message = (
                str(error).strip()
                or
                "Không thể xuất danh sách Alarm."
            )

        finally:

            if workbook is not None:

                try:
                    workbook.close()

                except Exception:
                    pass

        self.signals.finished.emit(
            self.output_path,
            error_message,
        )


# ============================================================
# TABLE MODEL
# ============================================================

class AlarmTableModel(
    QAbstractTableModel
):

    HEADERS = [
        "No",
        "Date",
        "Machine",
        "Slot Test",
        "FAIL Comment",
        "Start Time",
        "End Time",
        "Total Test",
        "PASS",
        "FAIL",
        "Yield",
        "Last 15",
        "Last 15 cùng Model",
        "Last 30",
        "Last 30 cùng Model",
        "SCRAP CODE",
        "Model",
        "Status",
        "Date Complete",
        "Quick Check Result",
        "CAL Check Result",
        "Engineer Action",
        "Monitor Result Day 1",
        "Monitor Result Day 2",
        "Monitor Result Day 3",
        "Comment",
    ]

    edit_requested = pyqtSignal(
        object,
        str,
        str,
    )

    def __init__(
        self,
        parent=None,
    ):

        super().__init__(
            parent
        )

        self.rows = []

        self.editing_enabled = True

    # ========================================================
    # ROW COUNT
    # ========================================================

    def rowCount(
        self,
        parent=QModelIndex(),
    ) -> int:

        if parent.isValid():
            return 0

        return len(
            self.rows
        )

    # ========================================================
    # COLUMN COUNT
    # ========================================================

    def columnCount(
        self,
        parent=QModelIndex(),
    ) -> int:

        if parent.isValid():
            return 0

        return len(
            self.HEADERS
        )

    # ========================================================
    # DATA
    # ========================================================

    def data(
        self,
        index,
        role=Qt.DisplayRole,
    ):

        if not index.isValid():
            return None

        if not (
            0 <= index.row()
            < len(self.rows)
        ):
            return None

        row = self.rows[
            index.row()
        ]

        # ====================================================
        # ALIGNMENT
        # ====================================================

        if role == Qt.TextAlignmentRole:

            return int(
                Qt.AlignCenter
            )

        # ====================================================
        # BACKGROUND
        # ====================================================

        if role == Qt.BackgroundRole:

            status = str(
                getattr(
                    row,
                    "status",
                    "",
                )
                or ""
            ).strip()

            # ------------------------------------------------
            # STATUS - TOÀN BỘ DÒNG
            # ------------------------------------------------

            if status == "Đã hoàn thành":

                return QColor(
                    "#DCFCE7"
                )

            if status == "Đang tiến hành":

                return QColor(
                    "#FEF3C7"
                )

            if status == "Chưa tiến hành":

                return QColor(
                    "#FEE2E2"
                )

            # ------------------------------------------------
            # QUICK / CAL
            # ------------------------------------------------

            if index.column() in (
                COL_QUICK_CHECK,
                COL_CAL_CHECK,
            ):

                field_name = (
                    "quick_check_result"
                    if index.column()
                    == COL_QUICK_CHECK
                    else "cal_check_result"
                )

                check_result = str(
                    getattr(
                        row,
                        field_name,
                        "",
                    )
                    or ""
                ).strip().upper()

                if check_result == "PASS":

                    return QColor(
                        "#DCFCE7"
                    )

                if check_result == "FAIL":

                    return QColor(
                        "#FEE2E2"
                    )

            # ------------------------------------------------
            # MONITOR
            # ------------------------------------------------

            if index.column() in (
                COL_MONITOR_DAY1,
                COL_MONITOR_DAY2,
                COL_MONITOR_DAY3,
            ):

                monitor_fields = {
                    COL_MONITOR_DAY1:
                        "monitor_day1",

                    COL_MONITOR_DAY2:
                        "monitor_day2",

                    COL_MONITOR_DAY3:
                        "monitor_day3",
                }

                monitor_value = str(
                    getattr(
                        row,
                        monitor_fields[
                            index.column()
                        ],
                        "",
                    )
                    or ""
                ).strip().upper()

                if monitor_value == "PASS":

                    return QColor(
                        "#DCFCE7"
                    )

                if monitor_value == "FAIL":

                    return QColor(
                        "#FEE2E2"
                    )

        # ====================================================
        # DISPLAY / EDIT
        # ====================================================

        if role not in (
            Qt.DisplayRole,
            Qt.EditRole,
        ):

            return None

        def value(
            name,
            default="",
        ):

            result = getattr(
                row,
                name,
                default,
            )

            if result is None:
                return default

            return result

        values = (
            index.row() + 1,
            value("alarm_date"),
            value("machine"),
            value("slot"),
            value("reason"),
            value("start_datetime"),
            value("end_datetime"),
            value("total_test", 0),
            value("pass_count", 0),
            value("fail_count", 0),
            self._format_percent(value("yield_percent")),
            self._format_percent(value("yield_15")),
            self._format_percent(value("yield_15_model")),
            self._format_percent(value("yield_30")),
            self._format_percent(value("yield_30_model")),
            value("scrap_codes"),
            value("models"),
            value(
                "status",
                "Chưa tiến hành",
            ),
            value("date_complete"),
            value("quick_check_result"),
            value("cal_check_result"),
            value("engineer_action"),
            value("monitor_day1"),
            value("monitor_day2"),
            value("monitor_day3"),
            value("comment"),
        )

        return values[
            index.column()
        ]

    @staticmethod
    def _format_percent(value):
        if value is None or value == "":
            return "-"
        try:
            return f"{float(value):.2f}%"
        except (TypeError, ValueError):
            return str(value)

    # ========================================================
    # FLAGS
    # ========================================================

    def flags(
        self,
        index,
    ):

        flags = super().flags(
            index
        )

        if (
            index.isValid()
            and index.column()
            in EDITABLE_COLUMNS
            and self.editing_enabled
        ):

            flags |= Qt.ItemIsEditable

        return flags

    # ========================================================
    # SET DATA
    # ========================================================

    def setData(
        self,
        index,
        value,
        role=Qt.EditRole,
    ):

        if (
            role != Qt.EditRole
            or not index.isValid()
            or index.column()
            not in EDITABLE_COLUMNS
            or not self.editing_enabled
        ):

            return False

        if not (
            0 <= index.row()
            < len(self.rows)
        ):

            return False

        row = self.rows[
            index.row()
        ]

        field_map = {
            COL_STATUS:
                "status",

            COL_QUICK_CHECK:
                "quick_check_result",

            COL_CAL_CHECK:
                "cal_check_result",

            COL_ENGINEER_ACTION:
                "engineer_action",

            COL_COMMENT:
                "comment",
        }

        field = field_map.get(
            index.column()
        )

        if not field:
            return False

        old_value = getattr(
            row,
            field,
            "",
        )

        if old_value is None:
            old_value = ""

        value = (
            str(value)
            if value is not None
            else ""
        )

        old_value = str(
            old_value
        )

        if value == old_value:
            return True

        self.edit_requested.emit(
            row,
            field,
            value,
        )

        return True

    # ========================================================
    # UPDATE ROW
    # ========================================================

    def update_row(
        self,
        latest,
    ):

        if latest is None:
            return

        latest_key = (
            self._row_key(
                latest
            )
        )

        for index, current in enumerate(
            self.rows
        ):

            if (
                self._row_key(
                    current
                )
                == latest_key
            ):

                self.rows[
                    index
                ] = latest

                self.dataChanged.emit(
                    self.index(
                        index,
                        0,
                    ),
                    self.index(
                        index,
                        self.columnCount()
                        - 1,
                    ),
                    [
                        Qt.DisplayRole,
                        Qt.EditRole,
                        Qt.BackgroundRole,
                    ],
                )

                return

    @staticmethod
    def _row_key(
        row,
    ):

        return (
            str(
                getattr(
                    row,
                    "alarm_type",
                    "",
                )
                or ""
            ),

            str(
                getattr(
                    row,
                    "alarm_date",
                    "",
                )
                or ""
            ),

            str(
                getattr(
                    row,
                    "machine",
                    "",
                )
                or ""
            ),

            str(
                getattr(
                    row,
                    "slot",
                    "",
                )
                or ""
            ),
        )

    # ========================================================
    # HEADER
    # ========================================================

    def headerData(
        self,
        section,
        orientation,
        role=Qt.DisplayRole,
    ):

        if (
            role == Qt.DisplayRole
            and orientation
            == Qt.Horizontal
            and 0 <= section
            < len(self.HEADERS)
        ):

            return self.HEADERS[
                section
            ]

        if role == Qt.TextAlignmentRole:

            return int(
                Qt.AlignCenter
            )

        return None

    # ========================================================
    # ROWS
    # ========================================================

    def set_rows(
        self,
        rows,
    ):

        self.beginResetModel()

        self.rows = list(
            rows or []
        )

        self.endResetModel()

    def clear(self):

        self.set_rows(
            []
        )

    def row_at(
        self,
        row_index: int,
    ):

        if (
            0 <= row_index
            < len(self.rows)
        ):

            return self.rows[
                row_index
            ]

        return None


# ============================================================
# FILTER PROXY MODEL
# ============================================================

class AlarmFilterProxyModel(
    QSortFilterProxyModel
):

    filter_changed = pyqtSignal()

    def __init__(
        self,
        parent=None,
    ):

        super().__init__(
            parent
        )

        self._column_filters = {}

        self.setDynamicSortFilter(
            True
        )

    # ========================================================
    # NORMALIZE
    # ========================================================

    @staticmethod
    def normalized_value(
        value,
    ) -> str:

        if value is None:
            return ""

        return str(
            value
        )

    # ========================================================
    # NATURAL SORT
    # ========================================================

    @staticmethod
    def natural_sort_key(
        value: str,
    ):

        parts = re.split(
            r"(\d+)",
            value.casefold(),
        )

        return tuple(
            (
                (0, int(part))
                if part.isdigit()
                else (1, part)
            )
            for part in parts
        )

    # ========================================================
    # FILTER
    # ========================================================

    def filterAcceptsRow(
        self,
        source_row,
        source_parent,
    ):

        return self._row_matches_filters(
            source_row,
            source_parent,
        )

    def _row_matches_filters(
        self,
        source_row: int,
        source_parent=QModelIndex(),
        ignored_column: int | None = None,
    ):

        source_model = (
            self.sourceModel()
        )

        if source_model is None:
            return False

        for (
            column,
            selected_values,
        ) in self._column_filters.items():

            if column == ignored_column:
                continue

            value = source_model.data(
                source_model.index(
                    source_row,
                    column,
                    source_parent,
                ),
                Qt.DisplayRole,
            )

            if (
                self.normalized_value(
                    value
                )
                not in selected_values
            ):

                return False

        return True

    # ========================================================
    # COLUMN VALUES
    # ========================================================

    def values_for_column(
        self,
        column: int,
    ) -> list[str]:

        source_model = (
            self.sourceModel()
        )

        if source_model is None:
            return []

        values = set()

        for row in range(
            source_model.rowCount()
        ):

            if not self._row_matches_filters(
                row,
                ignored_column=column,
            ):

                continue

            value = source_model.data(
                source_model.index(
                    row,
                    column,
                ),
                Qt.DisplayRole,
            )

            values.add(
                self.normalized_value(
                    value
                )
            )

        return sorted(
            values,
            key=self.natural_sort_key,
        )

    # ========================================================
    # SELECTED VALUES
    # ========================================================

    def selected_values_for_column(
        self,
        column: int,
    ) -> set[str]:

        all_values = set(
            self.values_for_column(
                column
            )
        )

        return set(
            self._column_filters.get(
                column,
                all_values,
            )
        )

    # ========================================================
    # SET FILTER
    # ========================================================

    def set_column_filter(
        self,
        column: int,
        selected_values: set[str],
    ):

        all_values = set(
            self.values_for_column(
                column
            )
        )

        selected_values = set(
            selected_values
        )

        if selected_values == all_values:

            self._column_filters.pop(
                column,
                None,
            )

        else:

            self._column_filters[
                column
            ] = selected_values

        self.invalidateFilter()

        self.headerDataChanged.emit(
            Qt.Horizontal,
            column,
            column,
        )

        self.filter_changed.emit()

    # ========================================================
    # CLEAR FILTER
    # ========================================================

    def clear_filters(
        self,
    ):

        if not self._column_filters:
            return

        changed_columns = list(
            self._column_filters
        )

        self._column_filters.clear()

        self.invalidateFilter()

        for column in changed_columns:

            self.headerDataChanged.emit(
                Qt.Horizontal,
                column,
                column,
            )

        self.filter_changed.emit()

    # ========================================================
    # HAS FILTER
    # ========================================================

    def has_active_filters(
        self,
    ):

        return bool(
            self._column_filters
        )

    def is_column_filtered(
        self,
        column: int,
    ):

        return (
            column
            in self._column_filters
        )

    # ========================================================
    # HEADER TOOLTIP
    # ========================================================

    def headerData(
        self,
        section,
        orientation,
        role=Qt.DisplayRole,
    ):

        value = super().headerData(
            section,
            orientation,
            role,
        )

        if (
            orientation
            == Qt.Horizontal
            and role
            == Qt.ToolTipRole
        ):

            if self.is_column_filtered(
                section
            ):

                return (
                    "Đang lọc cột này. "
                    "Bấm để thay đổi bộ lọc."
                )

            return (
                "Bấm để lọc cột này."
            )

        return value


# ============================================================
# FILTER HEADER
# ============================================================

class AlarmFilterHeader(
    QHeaderView
):

    def paintSection(
        self,
        painter,
        rect,
        logical_index,
    ):

        if not rect.isValid():
            return

        model = self.model()

        is_filtered = (
            isinstance(
                model,
                AlarmFilterProxyModel,
            )
            and model.is_column_filtered(
                logical_index
            )
        )

        painter.save()

        background = QColor(
            "#DBEAFE"
            if is_filtered
            else "#E2E8F0"
        )

        foreground = QColor(
            "#0D47A1"
            if is_filtered
            else "#1E293B"
        )

        painter.fillRect(
            rect,
            background,
        )

        painter.setPen(
            QPen(
                QColor("#CBD5E1"),
                1,
            )
        )

        painter.drawLine(
            rect.topRight(),
            rect.bottomRight(),
        )

        painter.drawLine(
            rect.bottomLeft(),
            rect.bottomRight(),
        )

        font = painter.font()
        font.setBold(True)

        painter.setFont(
            font
        )

        painter.setPen(
            foreground
        )

        header_text = model.headerData(
            logical_index,
            Qt.Horizontal,
            Qt.DisplayRole,
        )

        text_rect = rect.adjusted(
            4,
            0,
            -20,
            0,
        )

        if logical_index in (COL_TARGET15, COL_TARGET30):
            header_text = str(header_text or "").replace("cùng Model", "cùng\nModel")

        painter.drawText(
            text_rect,
            int(
                Qt.AlignCenter
                | Qt.AlignVCenter
                | Qt.TextWordWrap
            ),
            str(
                header_text or ""
            ),
        )

        icon_x = (
            rect.right() - 10
        )

        icon_y = (
            rect.center().y()
        )

        painter.setPen(
            Qt.NoPen
        )

        # ====================================================
        # FILTERED
        # ====================================================

        if is_filtered:

            painter.setBrush(
                QColor("#1565C0")
            )

            funnel = QPolygon([
                QPoint(
                    icon_x - 6,
                    icon_y - 5,
                ),
                QPoint(
                    icon_x + 6,
                    icon_y - 5,
                ),
                QPoint(
                    icon_x + 2,
                    icon_y,
                ),
                QPoint(
                    icon_x + 2,
                    icon_y + 5,
                ),
                QPoint(
                    icon_x - 2,
                    icon_y + 3,
                ),
                QPoint(
                    icon_x - 2,
                    icon_y,
                ),
            ])

            painter.drawPolygon(
                funnel
            )

            painter.fillRect(
                rect.left(),
                rect.bottom() - 2,
                rect.width(),
                3,
                QColor("#1976D2"),
            )

        # ====================================================
        # NOT FILTERED
        # ====================================================

        else:

            painter.setBrush(
                QColor("#334155")
            )

            arrow = QPolygon([
                QPoint(
                    icon_x - 5,
                    icon_y - 3,
                ),
                QPoint(
                    icon_x + 5,
                    icon_y - 3,
                ),
                QPoint(
                    icon_x,
                    icon_y + 3,
                ),
            ])

            painter.drawPolygon(
                arrow
            )

        painter.restore()


# ============================================================
# TEST HISTORY DIALOG
# ============================================================

class AlarmTestHistoryDialog(
    QDialog
):

    HEADERS = [
        "No",
        "Date",
        "Time",
        "Slot",
        "Result",
        "Part No / Model",
        "LOT NO",
        "Interface",
        "SCRAP CODE",
    ]

    def __init__(
        self,
        machine: str,
        slot: int,
        records,
        parent=None,
    ):

        super().__init__(
            parent
        )

        self.setWindowTitle(
            f"Test History - "
            f"{machine} / Slot {slot}"
        )

        self.resize(
            1250,
            620,
        )

        layout = QVBoxLayout(
            self
        )

        layout.setContentsMargins(
            12,
            12,
            12,
            12,
        )

        layout.setSpacing(
            8
        )

        title = QLabel(
            f"30 lần test gần nhất | "
            f"Machine: {machine} | "
            f"Slot: {slot}"
        )

        title.setStyleSheet(
            """
            font-size: 14px;
            font-weight: 600;
            color: #1E293B;
            """
        )

        layout.addWidget(
            title
        )

        status = QLabel(
            f"Tổng số bản ghi: "
            f"{len(records)}"
        )

        status.setStyleSheet(
            """
            font-size: 12px;
            color: #64748B;
            """
        )

        layout.addWidget(
            status
        )

        table = QTableWidget(
            self
        )

        table.setColumnCount(
            len(self.HEADERS)
        )

        table.setHorizontalHeaderLabels(
            self.HEADERS
        )

        table.setRowCount(
            len(records)
        )

        table.setEditTriggers(
            QAbstractItemView.NoEditTriggers
        )

        table.setSelectionBehavior(
            QAbstractItemView.SelectRows
        )

        table.setAlternatingRowColors(
            True
        )

        table.verticalHeader().setVisible(
            False
        )

        table.verticalHeader().setDefaultSectionSize(
            28
        )

        table.setStyleSheet(
            """
            QTableWidget {
                gridline-color: #CBD5E1;
                border: 1px solid #94A3B8;
                background-color: #FFFFFF;
                alternate-background-color: #F8FAFC;
                font-size: 12px;
            }

            QTableWidget::item:selected {
                color: #0F172A;
                background-color: #BFDBFE;
            }

            QHeaderView::section {
                color: #1E293B;
                background-color: #E2E8F0;
                border: none;
                border-right: 1px solid #CBD5E1;
                border-bottom: 1px solid #94A3B8;
                font-weight: 600;
            }
            """
        )

        widths = (
            55,
            95,
            90,
            65,
            85,
            180,
            140,
            140,
            140,
        )

        for col, width in enumerate(
            widths
        ):

            table.setColumnWidth(
                col,
                width,
            )

        for row_index, record in enumerate(
            records
        ):

            values = (
                row_index + 1,
                getattr(
                    record,
                    "date",
                    "",
                ),
                getattr(
                    record,
                    "time",
                    "",
                ),
                getattr(
                    record,
                    "slot",
                    "",
                ),
                getattr(
                    record,
                    "result",
                    "",
                ),
                getattr(
                    record,
                    "partno",
                    "",
                ),
                getattr(
                    record,
                    "lotno",
                    "",
                ),
                getattr(
                    record,
                    "interface",
                    "",
                ),
                getattr(
                    record,
                    "scrapcode",
                    "",
                ),
            )

            for col, value in enumerate(
                values
            ):

                item = QTableWidgetItem(
                    str(
                        value or ""
                    )
                )

                item.setTextAlignment(
                    Qt.AlignCenter
                )

                if col == 4:

                    result = str(
                        value or ""
                    ).upper()

                    if result == "PASS":

                        item.setBackground(
                            QColor("#DCFCE7")
                        )

                    elif result == "FAIL":

                        item.setBackground(
                            QColor("#FEE2E2")
                        )

                table.setItem(
                    row_index,
                    col,
                    item,
                )

        layout.addWidget(
            table,
            1,
        )

        close_button = QPushButton(
            "Close",
            self,
        )

        close_button.setMinimumHeight(
            30
        )

        close_button.clicked.connect(
            self.accept
        )

        layout.addWidget(
            close_button
        )


class _LogHistorySignals(QObject):
    loaded = pyqtSignal(object, object, object)
    failed = pyqtSignal(str)


class _LogHistoryWorker(QRunnable):
    """Read the clicked row's real log records without blocking the UI."""

    def __init__(self, database_path, log_root, row):
        super().__init__()
        self.database_path = Path(database_path)
        self.log_root = log_root
        self.row = row
        self.signals = _LogHistorySignals()

    def run(self):
        try:
            from repositories.machine_slot_yield_repository import MachineSlotYieldRepository
            from services.log_alarm_history import load_alarm_history
            config = MachineSlotYieldRepository(
                self.database_path.parent / 'machine_slot_targets.json')
            _, _, model15, model30 = config.load_all_targets()
            machine = str(self.row.machine).strip()
            slot = int(self.row.slot)
            windows = load_alarm_history(self.log_root, machine, slot,
                                         self.row, model15, model30)
            self.signals.loaded.emit(machine, slot, (self.row.alarm_date, windows))
        except Exception as error:
            self.signals.failed.emit(str(error).strip() or 'Không đọc được lịch sử log.')


# ============================================================
# ALARM TAB
# ============================================================

class AlarmTab(QWidget):

    def __init__(
        self,
        database_path: Union[str, Path],
        parent=None,
    ):

        super().__init__(
            parent
        )

        self.database_path = Path(
            database_path
        )

        self._export_date_from = ""
        self._export_date_to = ""

        self._alarm_loading = False

        self._exporting = False

        self._export_worker = None

        self._save_worker = None
        self._history_worker = None
        self._history_progress = None
        self._pending_changes = {}

        self._engineers = []

        self._engineer_error = ""

        self._build_ui()

    # ========================================================
    # UI
    # ========================================================

    def _build_ui(
        self,
    ):

        layout = QVBoxLayout(
            self
        )

        layout.setContentsMargins(
            10,
            10,
            10,
            10,
        )

        layout.setSpacing(
            10
        )

        # ====================================================
        # TITLE
        # ====================================================

        self.title_label = QLabel(
            "Alarm"
        )

        self.title_label.setStyleSheet(
            """
            QLabel {
                color: #1E293B;
                font-size: 14px;
                font-weight: 600;
            }
            """
        )

        # ====================================================
        # STATUS LABEL
        # ====================================================

        self.status_label = QLabel(
            "Chọn From, To và bấm Apply Filter."
        )

        self.status_label.setStyleSheet(
            """
            QLabel {
                color: #64748B;
                font-size: 12px;
            }
            """
        )

        # ====================================================
        # MODEL
        # ====================================================

        self.table_model = (
            AlarmTableModel(
                self
            )
        )

        self.filter_model = (
            AlarmFilterProxyModel(
                self
            )
        )

        self.filter_model.setSourceModel(
            self.table_model
        )

        # ====================================================
        # TABLE
        # ====================================================

        self.table = QTableView(
            self
        )

        self.filter_header = (
            AlarmFilterHeader(
                Qt.Horizontal,
                self.table,
            )
        )

        self.table.setHorizontalHeader(
            self.filter_header
        )

        self.table.setModel(
            self.filter_model
        )

        # ====================================================
        # DOUBLE CLICK
        # ====================================================

        self.table.doubleClicked.connect(
            self._handle_double_click
        )

        self.table.setAlternatingRowColors(
            False
        )

        self.table.setSelectionBehavior(
            QAbstractItemView.SelectRows
        )

        self.table.setSelectionMode(
            QAbstractItemView.SingleSelection
        )

        self.table.setEditTriggers(
            QAbstractItemView.DoubleClicked
            | QAbstractItemView.EditKeyPressed
            | QAbstractItemView.SelectedClicked
        )

        self.table.setSortingEnabled(
            False
        )

        self.table.verticalHeader().setVisible(
            False
        )

        self.table.verticalHeader().setDefaultSectionSize(
            30
        )

        # ====================================================
        # HEADER
        # ====================================================

        header = (
            self.table.horizontalHeader()
        )

        header.setFixedHeight(
            50
        )

        header.setSectionResizeMode(
            QHeaderView.Fixed
        )

        header.setSectionsClickable(
            True
        )

        header.setHighlightSections(
            False
        )

        header.sectionClicked.connect(
            self._open_column_filter
        )

        self.filter_model.filter_changed.connect(
            self._on_column_filter_changed
        )

        self.filter_model.filter_changed.connect(
            header.viewport().update
        )

        # ====================================================
        # COLUMN WIDTH
        # ====================================================

        widths = (
            45,  # No
            70,  # Date
            110,  # Machine
            70,  # Slot Test
            300,  # FAIL Comment
            125,  # Start Time
            125,  # End Time
            110,  # Total Test
            65,  # PASS
            65,  # FAIL
            75,  # Yield
            75,  # Last 15
            135,  # Last 15 cùng Model
            75,  # Last 30
            135,  # Last 30 cùng Model
            130,  # SCRAP CODE
            120,  # Model
            125,  # Status
            140,  # Date Complete
            135,  # Quick Check
            125,  # CAL Check
            125,  # Engineer Action
            140,  # Monitor Day 1
            140,  # Monitor Day 2
            140,  # Monitor Day 3
            100,  # Comment
        )

        for column, width in enumerate(
            widths
        ):

            self.table.setColumnWidth(
                column,
                width,
            )

        # ====================================================
        # STYLE
        # ====================================================

        self.table.setStyleSheet(
            """
            QTableView {
                gridline-color: #CBD5E1;
                border: 1px solid #94A3B8;
                background-color: #FFFFFF;
                font-size: 12px;
            }

            QTableView::item:selected {
                color: #0F172A;
            }

            QHeaderView::section {
                color: #1E293B;
                background-color: #E2E8F0;
                border: none;
                border-right: 1px solid #CBD5E1;
                border-bottom: 1px solid #94A3B8;
                font-weight: 600;
            }
            """
        )

        # ====================================================
        # EDIT DELEGATE
        # ====================================================

        delegate = AlarmChoiceDelegate(
            self
        )

        for editable_column in EDITABLE_COLUMNS:

            self.table.setItemDelegateForColumn(
                editable_column,
                delegate
            )

        # ====================================================
        # EDIT SIGNAL
        # ====================================================

        self.table_model.edit_requested.connect(
            self._save_tracking
        )

        # ====================================================
        # STATUS BAR
        # ====================================================

        status_layout = QHBoxLayout()

        status_layout.setContentsMargins(
            0,
            0,
            0,
            0,
        )

        status_layout.setSpacing(
            10
        )

        status_layout.addWidget(
            self.status_label
        )

        status_layout.addStretch(
            1
        )

        # ====================================================
        # SAVE
        # ====================================================

        self.save_button = QPushButton(
            "Save",
            self,
        )

        self.save_button.setMinimumWidth(
            100
        )

        self.save_button.setMinimumHeight(
            30
        )

        self.save_button.setStyleSheet(
            """
            QPushButton {
                background-color: #16A34A;
                color: white;
                border: none;
                border-radius: 4px;
                padding: 6px 14px;
                font-weight: 600;
            }

            QPushButton:hover {
                background-color: #15803D;
            }

            QPushButton:pressed {
                background-color: #166534;
            }

            QPushButton:disabled {
                background-color: #BBF7D0;
                color: #F0FDF4;
            }
            """
        )

        self.save_button.setEnabled(False)
        self.save_button.clicked.connect(
            self._save_pending_changes
        )

        status_layout.addWidget(
            self.save_button
        )

        # ====================================================
        # EXPORT
        # ====================================================

        self.export_button = QPushButton(
            "Export Excel",
            self,
        )

        self.export_button.setMinimumWidth(
            130
        )

        self.export_button.setMinimumHeight(
            30
        )

        self.export_button.setStyleSheet(
            """
            QPushButton {
                background-color: #1976D2;
                color: white;
                border: none;
                border-radius: 4px;
                padding: 6px 14px;
                font-weight: 600;
            }

            QPushButton:hover {
                background-color: #1565C0;
            }

            QPushButton:pressed {
                background-color: #0D47A1;
            }

            QPushButton:disabled {
                background-color: #90CAF9;
                color: #E3F2FD;
            }
            """
        )

        self.export_button.setEnabled(
            False
        )

        self.export_button.clicked.connect(
            self._export_alarm_excel
        )

        status_layout.addWidget(
            self.export_button
        )

        # ====================================================
        # ADD TO LAYOUT
        # ====================================================

        layout.addWidget(
            self.title_label
        )

        layout.addLayout(
            status_layout
        )

        layout.addWidget(
            self.table,
            1,
        )

    # ========================================================
    # DOUBLE CLICK HANDLER
    # ========================================================

    def _handle_double_click(
        self,
        proxy_index,
    ):

        if (
            not proxy_index.isValid()
            or self._alarm_loading
            or self.is_saving()
        ):
            return

        column = proxy_index.column()

        # ====================================================
        # SLOT
        # ====================================================

        if column == COL_SLOT:

            self._show_test_history(
                proxy_index
            )

            return

        # ====================================================
        # EDITABLE
        # ====================================================

        if column in EDITABLE_COLUMNS:

            self.table.edit(
                proxy_index
            )

            return

    # ========================================================
    # TEST HISTORY
    # ========================================================

    def _show_test_history(
        self,
        proxy_index,
    ):

        if (
            self._alarm_loading
            or self.is_saving()
            or not proxy_index.isValid()
        ):
            return

        if proxy_index.column() != COL_SLOT:
            return

        source_index = (
            self.filter_model.mapToSource(
                proxy_index
            )
        )

        if not source_index.isValid():
            return

        try:

            alarm_row = (
                self.table_model.rows[
                    source_index.row()
                ]
            )

        except (
            IndexError,
            TypeError,
        ):

            return

        machine = str(
            getattr(
                alarm_row,
                "machine",
                "",
            )
            or ""
        ).strip()

        slot = getattr(
            alarm_row,
            "slot",
            "",
        )

        if not machine or slot in (
            None,
            "",
        ):

            QMessageBox.information(
                self,
                "Test History",
                "Không xác định được Machine hoặc Slot của Alarm.",
            )

            return

        if self._history_worker is not None:
            return
        try:
            from repositories.machine_slot_yield_repository import MachineSlotYieldRepository
            from repositories.auto_import_scheduler_repository import AutoImportSchedulerRepository
            repository = MachineSlotYieldRepository(
                self.database_path.parent / 'machine_slot_targets.json')
            log_root = repository.load_log_folder()
            if not log_root:
                config = AutoImportSchedulerRepository(self.database_path).get_config()
                log_root = str(getattr(config, 'log_folder', '') or '').strip()
            if not log_root or not Path(log_root).is_dir():
                QMessageBox.warning(self, 'Test History',
                    'Chưa cấu hình Log Folder hợp lệ trong Machine Slot Yield hoặc File Management.')
                return
            self._history_progress = QProgressDialog(
                'Đang đọc lịch sử trực tiếp từ Log Folder...', '', 0, 0, self)
            self._history_progress.setCancelButton(None)
            self._history_progress.setWindowTitle('Test History')
            self._history_progress.setMinimumDuration(0)
            self._history_progress.show()
            worker = _LogHistoryWorker(self.database_path, log_root, alarm_row)
            self._history_worker = worker
            worker.signals.loaded.connect(self._on_log_history_loaded)
            worker.signals.failed.connect(self._on_log_history_failed)
            QThreadPool.globalInstance().start(worker)
        except Exception as error:
            self._finish_log_history()
            QMessageBox.warning(self, 'Test History', str(error))

    def _finish_log_history(self):
        if self._history_progress is not None:
            self._history_progress.close()
            self._history_progress = None
        self._history_worker = None

    def _on_log_history_loaded(self, machine, slot, payload):
        self._finish_log_history()
        alarm_date, windows = payload
        if not windows:
            QMessageBox.information(self, 'Alarm History',
                'Không có chuỗi vi phạm khớp với Alarm này trong Log Folder hiện tại.')
            return
        from ui.machine_alarm_history_dialog import MachineAlarmHistoryDialog
        MachineAlarmHistoryDialog(machine, slot, alarm_date, windows, self).exec_()

    def _on_log_history_failed(self, message):
        self._finish_log_history()
        QMessageBox.warning(self, 'Test History', message)

    def set_loading(
        self,
        date_from: str,
        date_to: str,
    ):

        self._alarm_loading = True

        self.table_model.editing_enabled = False

        self.export_button.setEnabled(
            False
        )

        self.title_label.setText(
            "Alarm | "
            f"{date_from} - {date_to}"
        )

        self.status_label.setText(
            "Đang tải dữ liệu Alarm..."
        )

    # ========================================================
    # FINISH LOADING STATE
    # ========================================================

    def _finish_loading_state(
        self,
    ):

        """
        Luôn gọi hàm này sau khi load Alarm.

        Mục đích:
        - Không để _alarm_loading bị kẹt True.
        - Cho phép Edit lại Status / Quick / CAL...
        - Cập nhật trạng thái Export.
        """

        self._alarm_loading = False

        self.table_model.editing_enabled = (
            not self.is_saving()
        )

        self.export_button.setEnabled(
            self.filter_model.rowCount() > 0
            and not self._exporting
            and not self.is_saving()
        )

        self.save_button.setEnabled(
            bool(self._pending_changes)
            and not self._alarm_loading
            and not self.is_saving()
        )

    # ========================================================
    # LOAD DATA
    # ========================================================

    def load_data(
        self,
        result,
    ):

        """
        Load dữ liệu Alarm.

        QUAN TRỌNG:
        Toàn bộ hàm được bảo vệ bằng try/except/finally.

        Nếu EXE ở PC khác gặp lỗi:
        - Engineer config
        - database
        - dữ liệu Alarm
        - object result
        - bất kỳ exception nào khác

        thì _alarm_loading vẫn được reset.
        """

        try:

            # =================================================
            # RESULT NULL
            # =================================================

            if result is None:

                self.show_error(
                    "Không nhận được dữ liệu Alarm."
                )

                return

            # =================================================
            # DATE
            # =================================================

            date_from = str(
                getattr(
                    result,
                    "date_from",
                    "",
                )
                or ""
            )

            date_to = str(
                getattr(
                    result,
                    "date_to",
                    "",
                )
                or ""
            )

            self.title_label.setText(
                "Alarm | "
                f"{date_from} - {date_to}"
            )

            # =================================================
            # LOAD ENGINEERS
            # =================================================

            try:

                self._engineers = (
                    load_engineers(
                        self.database_path
                    )
                )

                self._engineer_error = ""

            except Exception as error:

                self._engineers = []

                self._engineer_error = (
                    str(error).strip()
                    or
                    "Không thể tải cấu hình Engineer."
                )

                # ---------------------------------------------
                # QUAN TRỌNG:
                # Không QMessageBox ở đây.
                #
                # Nếu chạy EXE trên PC khác mà thiếu config
                # Engineer thì không được làm gián đoạn Alarm.
                # ---------------------------------------------

                print(
                    "Alarm Engineer load error:",
                    error,
                )

            # =================================================
            # ONLY MACHINE SLOT YIELD
            # =================================================

            machine_slot_yield_rows = []

            rows = getattr(
                result,
                "rows",
                [],
            )

            if rows is None:
                rows = []

            for row in rows:

                try:

                    alarm_type = str(
                        getattr(
                            row,
                            "alarm_type",
                            "",
                        )
                        or ""
                    ).strip()

                except Exception:

                    alarm_type = ""

                if alarm_type.strip().upper() in ("MACHINE SLOT YIELD", "MACHINE_SLOT_YIELD"):

                    machine_slot_yield_rows.append(
                        row
                    )

            # =================================================
            # CLEAR FILTER
            # =================================================

            self.filter_model.clear_filters()

            # =================================================
            # LOAD TABLE
            # =================================================

            self._pending_changes.clear()

            self.table_model.set_rows(
                machine_slot_yield_rows
            )

            # =================================================
            # EXPORT DATE
            # =================================================

            self._export_date_from = (
                date_from
            )

            self._export_date_to = (
                date_to
            )

            # =================================================
            # STATUS
            # =================================================

            total = len(
                machine_slot_yield_rows
            )

            if total:

                if self._engineer_error:

                    self.status_label.setText(
                        f"Đã tải {total:,} Alarm. "
                        "Cảnh báo: không tải được Engineer."
                    )

                else:

                    self.status_label.setText(
                        f"Đã tải {total:,} Alarm."
                    )

            else:

                if self._engineer_error:

                    self.status_label.setText(
                        "Không có Alarm. "
                        "Cảnh báo: không tải được Engineer."
                    )

                else:

                    self.status_label.setText(
                        "Không có Alarm trong "
                        "khoảng thời gian đã chọn."
                    )

        except Exception as error:

            # =================================================
            # LOAD ERROR
            # =================================================

            print(
                "Alarm load_data error:",
                error,
            )

            self.table_model.clear()

            self._export_date_from = ""
            self._export_date_to = ""

            self.export_button.setEnabled(
                False
            )

            self.status_label.setText(
                "Lỗi khi tải dữ liệu Alarm."
            )

            try:

                QMessageBox.warning(
                    self,
                    "Alarm",
                    "Có lỗi khi tải dữ liệu Alarm:\n"
                    f"{error}",
                )

            except Exception:

                pass

        finally:

            # =================================================
            # QUAN TRỌNG NHẤT
            # =================================================
            #
            # Dù load thành công hay lỗi,
            # luôn mở khóa Alarm.
            #
            self._finish_loading_state()

    # ========================================================
    # ERROR
    # ========================================================

    def show_error(
        self,
        error_message: str,
    ):

        self.filter_model.clear_filters()

        self.table_model.clear()

        self._alarm_loading = False

        self.table_model.editing_enabled = True

        self._export_date_from = ""

        self._export_date_to = ""

        self.export_button.setEnabled(
            False
        )

        self.status_label.setText(
            "Không thể tải dữ liệu Alarm: "
            f"{error_message}"
        )

    # ========================================================
    # INVALIDATE
    # ========================================================

    def invalidate_cache(
        self,
    ):

        self._pending_changes.clear()
        self.save_button.setEnabled(False)

        self.filter_model.clear_filters()

        self.table_model.clear()

        self._alarm_loading = False

        self.table_model.editing_enabled = True

        self._export_date_from = ""

        self._export_date_to = ""

        self.export_button.setEnabled(
            False
        )

        self.status_label.setText(
            "Dữ liệu Alarm đã thay đổi. "
            "Vui lòng bấm Apply Filter."
        )

    # ========================================================
    # SAVE
    # ========================================================

    def is_saving(
        self,
    ):
        return self._save_worker is not None

    def has_unsaved_changes(self) -> bool:
        return bool(self._pending_changes)

    def _save_tracking(
        self,
        expected,
        field,
        value,
    ):
        """
        Chỉ cập nhật dữ liệu trên giao diện và đưa vào hàng chờ.
        Người dùng bấm Save mới ghi xuống DATABASE.
        """

        if self._alarm_loading or self.is_saving():
            return

        value = str(value or "").strip()

        if field == "status":
            if value not in (
                "Chưa tiến hành",
                "Đang tiến hành",
                "Đã hoàn thành",
            ):
                QMessageBox.warning(
                    self,
                    "Trạng thái không hợp lệ",
                    "Trạng thái chỉ được chọn:\n"
                    "- Chưa tiến hành\n"
                    "- Đang tiến hành\n"
                    "- Đã hoàn thành",
                )
                return

        if field in (
            "quick_check_result",
            "cal_check_result",
        ):
            value = value.upper()
            if value and value not in ("PASS", "FAIL"):
                QMessageBox.warning(
                    self,
                    "Dữ liệu không hợp lệ",
                    f"{field} chỉ được chọn PASS hoặc FAIL.",
                )
                return

        key = self.table_model._row_key(expected)

        pending = self._pending_changes.get(key)

        if pending is None:
            pending = {
                "expected": expected,
                "changes": {},
            }
            self._pending_changes[key] = pending

        pending["changes"][field] = value

        # Cập nhật ngay trên bảng để người dùng thấy dữ liệu vừa sửa.
        current = next(
            (
                row
                for row in self.table_model.rows
                if self.table_model._row_key(row) == key
            ),
            None,
        )

        if current is not None:
            updated = current

            if field == "status":
                updated = replace(
                    updated,
                    status=value,
                    date_complete=(
                        datetime.now().strftime(
                            "%d/%m/%Y %H:%M:%S"
                        )
                        if value == "Đã hoàn thành"
                        else ""
                    ),
                )
            else:
                updated = replace(
                    updated,
                    **{field: value},
                )

            self.table_model.update_row(updated)

        self.save_button.setEnabled(True)

        self.status_label.setText(
            f"Đã thay đổi {len(self._pending_changes):,} Alarm. "
            "Bấm Save để lưu."
        )

    def _save_pending_changes(self) -> None:
        """Ghi toàn bộ thay đổi đang chờ xuống database."""

        if self.is_saving():
            return

        if not self._pending_changes:
            self.status_label.setText(
                "Không có thay đổi cần lưu."
            )
            self.save_button.setEnabled(False)
            return

        changes = [
            (
                item["expected"],
                dict(item["changes"]),
            )
            for item in self._pending_changes.values()
            if item.get("changes")
        ]

        if not changes:
            self._pending_changes.clear()
            self.save_button.setEnabled(False)
            return

        answer = QMessageBox.question(
            self,
            "Xác nhận Save",
            (
                f"Lưu {len(changes):,} Alarm đã thay đổi "
                "vào database?"
            ),
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.Yes,
        )

        if answer != QMessageBox.Yes:
            return

        self.table_model.editing_enabled = False
        self.save_button.setEnabled(False)
        self.export_button.setEnabled(False)
        self.status_label.setText(
            f"Đang lưu {len(changes):,} Alarm..."
        )

        try:
            self._save_worker = AlarmTrackingSaveWorker(
                self.database_path,
                changes,
            )

            self._save_worker.signals.succeeded.connect(
                self._on_batch_save_succeeded
            )

            self._save_worker.signals.failed.connect(
                self._on_batch_save_failed
            )

            self._save_worker.signals.finished.connect(
                self._on_batch_save_finished
            )

            QThreadPool.globalInstance().start(
                self._save_worker
            )

        except BaseException as error:
            self._save_worker = None
            self.table_model.editing_enabled = True
            self.save_button.setEnabled(
                bool(self._pending_changes)
            )
            self.export_button.setEnabled(
                self.filter_model.rowCount() > 0
            )
            QMessageBox.critical(
                self,
                "Không thể lưu Alarm",
                str(error).strip()
                or type(error).__name__,
            )

    def _on_batch_save_succeeded(self, latest_rows) -> None:
        for latest in latest_rows or []:
            self.table_model.update_row(latest)

        self._pending_changes.clear()

        self.status_label.setText(
            f"Đã lưu {len(latest_rows or []):,} Alarm thành công."
        )

        # Dashboard cần đọc lại dữ liệu mới.
        parent = self.parent()
        if parent is not None:
            try:
                parent.dashboard_tab.invalidate_cache()
            except Exception:
                pass

    def _on_batch_save_failed(self, error_message: str) -> None:
        self.status_label.setText(
            "Lưu Alarm thất bại."
        )
        QMessageBox.warning(
            self,
            "Không thể lưu Alarm",
            error_message,
        )

    def _on_batch_save_finished(self) -> None:
        self._save_worker = None
        self.table_model.editing_enabled = (
            not self._alarm_loading
        )
        self.save_button.setEnabled(
            bool(self._pending_changes)
            and not self._alarm_loading
        )
        self.export_button.setEnabled(
            self.filter_model.rowCount() > 0
            and not self._exporting
            and not self._alarm_loading
        )

    # ========================================================
    # COLUMN FILTER
    # ========================================================

    def _open_column_filter(
        self,
        column: int,
    ):

        if (
            self._alarm_loading
            or self.is_saving()
        ):

            return

        raw_values = (
            self.filter_model.values_for_column(
                column
            )
        )

        if not raw_values:
            return

        try:

            from ui.filter_value_dialog import (
                FilterValueDialog,
            )

        except ImportError as error:

            QMessageBox.warning(
                self,
                "Filter",
                f"Không thể mở bộ lọc cột:\n{error}",
            )

            return

        blank_label = "(Trống)"

        display_by_raw = {
            value:
                blank_label
                if value == ""
                else value
            for value in raw_values
        }

        raw_by_display = {
            display: raw
            for raw, display
            in display_by_raw.items()
        }

        selected_raw = (
            self.filter_model
            .selected_values_for_column(
                column
            )
        )

        dialog = FilterValueDialog(
            title=(
                f"Filter "
                f"{AlarmTableModel.HEADERS[column]}"
            ),

            values=[
                display_by_raw[value]
                for value in raw_values
            ],

            selected_values={
                display_by_raw[value]
                for value in selected_raw
                if value in display_by_raw
            },

            apply_immediately=True,

            compact=True,

            parent=self,
        )

        def apply_selected_values(
            selected_display_values,
        ):

            self.filter_model.set_column_filter(
                column,
                {
                    raw_by_display[value]
                    for value in selected_display_values
                    if value in raw_by_display
                },
            )

        dialog.selection_changed.connect(
            apply_selected_values
        )

        self._position_column_filter_dialog(
            dialog,
            column,
        )

        dialog.exec_()

    # ========================================================
    # FILTER POSITION
    # ========================================================

    def _position_column_filter_dialog(
        self,
        dialog,
        column: int,
    ):

        header = (
            self.table.horizontalHeader()
        )

        column_x = (
            header.sectionViewportPosition(
                column
            )
        )

        anchor = (
            header.viewport()
            .mapToGlobal(
                QPoint(
                    column_x,
                    header.height(),
                )
            )
        )

        screen = QApplication.screenAt(
            anchor
        )

        if screen is None:

            screen = (
                QApplication.primaryScreen()
            )

        target_x = anchor.x()

        target_y = anchor.y()

        if screen is not None:

            available = (
                screen.availableGeometry()
            )

            target_x = max(
                available.left(),
                min(
                    target_x,
                    available.right()
                    - dialog.width()
                    + 1,
                ),
            )

            if (
                target_y
                + dialog.height()
                > available.bottom()
                + 1
            ):

                target_y = max(
                    available.top(),
                    anchor.y()
                    - header.height()
                    - dialog.height(),
                )

        dialog.move(
            target_x,
            target_y,
        )

    # ========================================================
    # FILTER CHANGED
    # ========================================================

    def _on_column_filter_changed(
        self,
    ):

        visible_count = (
            self.filter_model.rowCount()
        )

        total_count = (
            self.table_model.rowCount()
        )

        self.export_button.setEnabled(
            visible_count > 0
            and not self._alarm_loading
            and not self._exporting
            and not self.is_saving()
        )

        if (
            self.filter_model
            .has_active_filters()
        ):

            self.status_label.setText(
                f"Đang hiển thị: "
                f"{visible_count:,} / "
                f"Tổng Alarm: "
                f"{total_count:,}"
            )

        elif total_count:

            self.status_label.setText(
                f"Tổng Alarm: "
                f"{total_count:,}"
            )

        else:

            self.status_label.setText(
                "Không có Alarm."
            )

    # ========================================================
    # EXPORT
    # ========================================================

    def _export_alarm_excel(
        self,
    ):

        if (
            self._exporting
            or self._alarm_loading
            or self.is_saving()
        ):

            return

        model = self.table.model()

        if (
            model is None
            or model.rowCount() == 0
        ):

            QMessageBox.information(
                self,
                "Export Excel",
                "Không có Alarm để xuất.",
            )

            return

        default_name = (
            f"Machine Slot Yield Alarm "
            f"{self._export_date_from}"
            f" - "
            f"{self._export_date_to}.xlsx"
        )

        headers = list(
            AlarmTableModel.HEADERS
        )

        rows = []

        for row_index in range(
            model.rowCount()
        ):

            row_values = []

            for column in range(
                model.columnCount()
            ):

                row_values.append(
                    model.data(
                        model.index(
                            row_index,
                            column,
                        ),
                        Qt.DisplayRole,
                    )
                )

            rows.append(
                row_values
            )

        output_path, _ = (
            QFileDialog.getSaveFileName(
                self,
                "Export Machine Slot Yield Alarm",
                default_name,
                "Excel Workbook (*.xlsx)",
            )
        )

        if not output_path:
            return

        if not output_path.lower().endswith(
            ".xlsx"
        ):

            output_path += ".xlsx"

        if Path(
            output_path
        ).exists():

            answer = QMessageBox.question(
                self,
                "Export Excel",
                "File đã tồn tại. "
                "Bạn có muốn ghi đè không?",

                QMessageBox.Yes
                | QMessageBox.No,

                QMessageBox.No,
            )

            if answer != QMessageBox.Yes:
                return

        # ====================================================
        # START EXPORT
        # ====================================================

        self._exporting = True

        self.export_button.setEnabled(
            False
        )

        self.export_button.setText(
            "Exporting..."
        )

        self._export_worker = (
            AlarmExportWorker(
                output_path=output_path,
                headers=headers,
                rows=rows,
            )
        )

        self._export_worker.signals.finished.connect(
            self._on_alarm_export_finished
        )

        QThreadPool.globalInstance().start(
            self._export_worker
        )

    # ========================================================
    # EXPORT FINISHED
    # ========================================================

    def _on_alarm_export_finished(
        self,
        output_path: str,
        error_message: str,
    ):

        self._exporting = False

        self._export_worker = None

        self.export_button.setText(
            "Export Excel"
        )

        self.export_button.setEnabled(
            not self._alarm_loading
            and not self.is_saving()
            and self.filter_model.rowCount()
            > 0
        )

        if error_message:

            QMessageBox.warning(
                self,
                "Export Excel",
                "Xuất Excel thất bại:\n"
                f"{error_message}",
            )

            return

        QMessageBox.information(
            self,
            "Export Excel",
            "Đã xuất danh sách Machine Slot Yield Alarm:\n"
            f"{output_path}",
        )